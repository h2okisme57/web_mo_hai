import io
import re
import time
import zipfile
import xml.etree.ElementTree as ET
import streamlit as st
from google import genai
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt

# Danh sách model dự phòng theo thứ tự ưu tiên tránh lỗi 429 và 503
MODEL_CANDIDATES = [
    "gemini-2.0-flash",
    "gemini-2.0-flash-lite",
    "gemini-1.5-flash",
    "gemini-2.5-flash",
    "gemini-2.0-flash-exp",
]

st.set_page_config(
    page_title="Trợ lý Giảng dạy & Slide 7991 - Cô Nguyễn Thị Ngọc Yến",
    page_icon="📚",
    layout="wide"
)

# Banner định danh giao diện
st.markdown("## 📚 HỆ THỐNG TRỢ LÝ GIẢNG DẠY SỐ & BIÊN SOẠN SLIDE")
st.markdown("**Giáo viên thực hiện:** Nguyễn Thị Ngọc Yến (Môn: Ngữ văn 6 – Trường THCS Nguyễn Đáng)")
st.caption("Chương trình chuyển giao kỹ thuật ứng dụng AI trong GD – Tỉnh Vĩnh Long (10/2026)")
st.divider()

# Sidebar: Thiết lập hệ thống & Thông tin bài dạy
with st.sidebar:
    st.header("⚙️ Thiết lập hệ thống")
    api_key_input = st.text_input("Nhập Google Gemini API Key:", type="password")
    api_key = api_key_input.strip() or st.secrets.get("GEMINI_API_KEY", "").strip()
    
    # Nút kiểm tra API Key tự động dò model khả dụng
    if st.button("🔍 Kiểm tra trạng thái AI", use_container_width=True):
        if not api_key:
            st.error("Chưa nhập API Key!")
        else:
            with st.spinner("Đang tìm model còn quota khả dụng..."):
                client = genai.Client(api_key=api_key)
                connected_model = None
                last_err = None
                
                for candidate in MODEL_CANDIDATES:
                    try:
                        res = client.models.generate_content(
                            model=candidate,
                            contents="ping"
                        )
                        if res:
                            connected_model = candidate
                            break
                    except Exception as err:
                        last_err = err
                        continue
                
                if connected_model:
                    st.session_state["active_model"] = connected_model
                    st.success(f"✅ Kết nối thành công!\n\nĐang sử dụng: `{connected_model}`")
                else:
                    st.error(f"❌ Tất cả model trong danh sách đều hết quota hoặc lỗi: {last_err}")

    st.divider()
    st.header("📋 Thông tin bài dạy thực nghiệm")
    gv_name = st.text_input("Họ tên giáo viên:", value="Nguyễn Thị Ngọc Yến")
    mon_hoc = st.selectbox("Môn học phụ trách:", ["Ngữ văn", "Tiếng Anh", "Lịch sử và Địa lí", "GDCD"], index=0)
    khoi_lop = st.selectbox("Khối lớp:", [f"Lớp {i}" for i in range(6, 13)], index=0)
    ten_bai_hoc = st.text_input("Tên bài học / Văn bản:", value="Bài học đường đời đầu tiên (Tô Hoài)")
    thoi_luong = st.number_input("Thời lượng (tiết):", min_value=1, max_value=6, value=3)
    yeu_cau_can_dat = st.text_area(
        "Mục tiêu cần đạt (CT GDPT 2018):",
        height=100,
        value="- Nhận biết được một số yếu tố của truyện đồng thoại: cốt truyện, nhân vật, lời người kể chuyện ngôi thứ nhất.\n- Phân tích được đặc điểm nhân vật Dế Mèn và Dế Choắt qua ngoại hình, hành động, ngôn ngữ.\n- Rút ra bài học về thái độ sống khiêm nhường, vị tha và cách ứng xử với bạn bè."
    )

# Hàm đọc file docx an toàn có fallback giải nén XML
def read_uploaded_file(uploaded_file):
    try:
        raw_bytes = uploaded_file.read()
        if uploaded_file.name.endswith(".docx"):
            # Cách 1: Thử bằng python-docx chuẩn
            try:
                doc = Document(io.BytesIO(raw_bytes))
                full_text = []
                for p in doc.paragraphs:
                    if p.text.strip():
                        full_text.append(p.text.strip())
                for table in doc.tables:
                    for row in table.rows:
                        row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_text:
                            full_text.append(" | ".join(row_text))
                extracted = "\n".join(full_text)
                if extracted.strip():
                    return extracted
            except Exception:
                pass

            # Cách 2: Mở trực tiếp file ZIP XML nếu header docx bị lỗi định dạng
            try:
                with zipfile.ZipFile(io.BytesIO(raw_bytes)) as z:
                    target_xml = None
                    for name in z.namelist():
                        if name.endswith("word/document.xml") or name.endswith("document.xml"):
                            target_xml = name
                            break
                    if target_xml:
                        xml_content = z.read(target_xml)
                        tree = ET.fromstring(xml_content)
                        texts = [elem.text for elem in tree.iter() if elem.tag.endswith('t') and elem.text]
                        return "\n".join(texts)
            except Exception:
                pass

        elif uploaded_file.name.endswith(".txt"):
            return raw_bytes.decode("utf-8", errors="ignore")

    except Exception as e:
        st.error(f"Lỗi khi đọc file tài liệu: {e}")
        return ""
    return ""

# Hàm gọi Gemini AI tự động né lỗi 429 và 503
def call_gemini(prompt_text, key):
    if not key:
        st.error("Vui lòng nhập API Key ở thanh bên trái!")
        return None
    try:
        client = genai.Client(api_key=key.strip())
        last_error = None
        
        active_m = st.session_state.get("active_model")
        trial_models = [active_m] + [m for m in MODEL_CANDIDATES if m != active_m] if active_m else MODEL_CANDIDATES

        for model_name in trial_models:
            for retry in range(2):
                try:
                    response = client.models.generate_content(
                        model=model_name,
                        contents=prompt_text,
                    )
                    
                    if response and hasattr(response, "text") and response.text:
                        st.session_state["active_model"] = model_name
                        return response.text
                    
                    if response and response.candidates:
                        parts_text = "".join([p.text for p in response.candidates[0].content.parts if hasattr(p, "text") and p.text])
                        if parts_text.strip():
                            st.session_state["active_model"] = model_name
                            return parts_text
                except Exception as err:
                    last_error = err
                    err_msg = str(err).lower()
                    # Bỏ qua ngay lập tức nếu hết quota hoặc model lỗi
                    if "429" in err_msg or "resource_exhausted" in err_msg or "quota" in err_msg or "503" in err_msg or "404" in err_msg:
                        break
                    else:
                        time.sleep(1)
                        continue

        st.error(f"Lỗi kết nối API trên tất cả mô hình: {last_error}")
        return None
    except Exception as e:
        st.error(f"Lỗi khởi tạo client: {e}")
        return None

# Xuất file Word (.docx)
def export_docx(title, content, author="Nguyễn Thị Ngọc Yến"):
    doc = Document()
    doc.add_heading(title, level=0)
    p = doc.add_paragraph()
    p.add_run(f"Giáo viên thực hiện: {author}\n").bold = True
    p.add_run("Quy chuẩn chuyên môn: Công văn 5512/BGDĐT & Công văn 7991/BGDĐT-GDTrH\n").italic = True
    for para in content.split("\n"):
        doc.add_paragraph(para)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

# Hàm xuất PowerPoint chuẩn Bài tập 3 (TO - VỪA - NHỎ + Speaker Notes)
def export_advanced_pptx(title, raw_text, author):
    prs = Presentation()
    prs.slide_width = Inches(13.333) # 16:9 widescreen
    prs.slide_height = Inches(7.5)
    
    # 1. Slide bìa
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = re.sub(r'[*#_]', '', title).strip()
    slide.placeholders[1].text = f"Môn học: {mon_hoc} - {khoi_lop}\nGiáo viên biên soạn: {author}"
    
    # 2. Bóc tách từng slide bằng Regex
    slide_blocks = re.split(r'(?:#{1,4}\s*)?(?:\*\*)?SLIDE\s*\d*[:\-–]?', raw_text, flags=re.IGNORECASE)
    
    for block in slide_blocks:
        block = block.strip()
        if not block:
            continue
            
        if ("dưới đây là kịch bản" in block.lower() or "powerpoint" in block.lower()) and len(block) < 180:
            if not any(k in block.upper() for k in ["TO:", "VỪA:", "NHỎ:"]):
                continue

        lines = [l.strip() for l in block.split("\n") if l.strip()]
        if not lines:
            continue
            
        slide_title = re.sub(r'[*#_]', '', lines[0]).strip()
        body_lines = []
        speaker_notes = []
        is_notes = False
        
        for l in lines[1:]:
            clean_l = re.sub(r'[*#_]', '', l).strip()
            if not clean_l:
                continue
            if any(k in clean_l.upper() for k in ["LỜI DẪN GV:", "GHI CHÚ GV:", "NOTES:"]):
                is_notes = True
                clean_l = re.sub(r'^(LỜI DẪN GV|GHI CHÚ GV|NOTES)[:\-–]\s*', '', clean_l, flags=re.IGNORECASE)
                if clean_l:
                    speaker_notes.append(clean_l)
                continue
                
            if is_notes:
                speaker_notes.append(clean_l)
            else:
                body_lines.append(clean_l)
                
        # Tạo Slide nội dung
        new_slide = prs.slides.add_slide(prs.slide_layouts[1])
        new_slide.shapes.title.text = slide_title
        
        tf = new_slide.placeholders[1].text_frame
        tf.clear()
        
        for bl in body_lines:
            bl_clean = bl.lstrip("-•* ")
            if not bl_clean:
                continue
            p = tf.add_paragraph()
            
            # Phân cấp kích thước chữ TO - VỪA - NHỎ
            if bl_clean.upper().startswith("TO:"):
                p.text = re.sub(r'^TO[:\-–]\s*', '', bl_clean, flags=re.IGNORECASE)
                p.font.size = Pt(24)
                p.font.bold = True
            elif bl_clean.upper().startswith("VỪA:"):
                p.text = re.sub(r'^VỪA[:\-–]\s*', '', bl_clean, flags=re.IGNORECASE)
                p.font.size = Pt(19)
            elif bl_clean.upper().startswith("NHỎ:"):
                p.text = re.sub(r'^NHỎ[:\-–]\s*', '', bl_clean, flags=re.IGNORECASE)
                p.font.size = Pt(15)
                p.font.italic = True
            else:
                p.text = bl_clean
                p.font.size = Pt(19)
                
        # Ghi lời dẫn giáo viên vào Speaker Notes
        if speaker_notes:
            notes_slide = new_slide.notes_slide
            text_frame = notes_slide.notes_text_frame
            text_frame.text = "\n".join(speaker_notes)
            
    bio = io.BytesIO()
    prs.save(bio)
    return bio.getvalue()

# Bốn phân hệ điều hướng
tab1, tab2, tab3, tab4 = st.tabs([
    "📖 1. Kế hoạch bài dạy (Bài 1)", 
    "📝 2. Đề kiểm tra CV 7991 (Bài 1)", 
    "💡 3. Trợ lý SKKN (Bài 2)", 
    "🖥️ 4. Nạp KHBD xuất Slide PPTX (Bài 3)"
])

# TAB 1: KẾ HOẠCH BÀI DẠY (BÀI TẬP 1)
with tab1:
    st.subheader("📋 Thiết kế Kế hoạch bài dạy chuẩn CV 5512")
    if st.button("🚀 Khởi tạo Kế hoạch bài dạy", type="primary", key="btn_run_khbd"):
        with st.spinner("AI đang soạn tiến trình dạy học theo chuẩn 5512..."):
            prompt = f"""
            Đóng vai trò là cố vấn phương pháp dạy học Ngữ văn cấp THCS. Hãy soạn Kế hoạch bài dạy hoàn chỉnh cho:
            - Môn: {mon_hoc} - {khoi_lop}
            - Bài dạy: {ten_bai_hoc} (Thời lượng: {thoi_luong} tiết)
            - Yêu cầu cần đạt: {yeu_cau_can_dat}
            - Giáo viên thực hiện: {gv_name}

            Cấu trúc kế hoạch bài dạy:
            I. MỤC TIÊU (Năng lực đặc thù, Năng lực chung, Phẩm chất nhân ái).
            II. THIẾT BỊ DẠY HỌC & HỌC LIỆU
            III. TIẾN TRÌNH DẠY HỌC (Gồm 4 hoạt động: Khởi động, Hình thành kiến thức, Luyện tập, Vận dụng; mỗi hoạt động đủ 4 bước: Chuyển giao -> Thực hiện -> Báo cáo -> Nhận định).
            Ký tên người soạn: {gv_name} - GV.
            """
            st.session_state["out_khbd"] = call_gemini(prompt, api_key)
            
    if st.session_state.get("out_khbd"):
        st.markdown(st.session_state["out_khbd"])
        f_docx = export_docx(f"KHBD_{ten_bai_hoc}", st.session_state["out_khbd"], gv_name)
        st.download_button(
            "📥 Tải về file Word KHBD (.docx)", 
            data=f_docx, 
            file_name=f"KHBD_{ten_bai_hoc}_{gv_name}.docx", 
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

# TAB 2: ĐỀ KIỂM TRA ĐỊNH KỲ THEO CÔNG VĂN 7991 (BÀI TẬP 1)
with tab2:
    st.subheader("📝 Ma trận, Bản đặc tả & Đề kiểm tra định kỳ (Chuẩn CV 7991)")
    c1, c2 = st.columns(2)
    with c1:
        exam_time = st.selectbox("Thời lượng kiểm tra:", ["45 phút", "60 phút", "90 phút"])
    with c2:
        has_short = st.checkbox("Môn có câu hỏi Trả lời ngắn", value=False)
        
    if st.button("🚀 Khởi tạo Bộ hồ sơ kiểm tra 7991", type="primary", key="btn_run_exam"):
        with st.spinner("Đang lập ma trận và đặc tả theo Công văn 7991/BGDĐT-GDTrH..."):
            prompt_ex = f"""
            Xây dựng bộ hồ sơ kiểm tra định kỳ chuẩn xác theo CÔNG VĂN 7991/BGDĐT-GDTrH:
            - Môn học: {mon_hoc} - {khoi_lop} | Bài: {ten_bai_hoc} | Thời gian: {exam_time}
            - Giáo viên thẩm định / ra đề: {gv_name}

            QUY CHUẨN CÔNG VĂN 7991:
            1. Tỉ lệ điểm:
               - Phần I. Trắc nghiệm Nhiều lựa chọn: 3,0 điểm (30%).
               - Phần II. Trắc nghiệm Đúng - Sai (mỗi câu gồm 4 lệnh hỏi a, b, c, d): 2,0 điểm (20%).
               {"- Phần III. Trắc nghiệm Trả lời ngắn: 2,0 điểm (20%)." if has_short else "- Phần III: Chuyển 2.0 điểm của Trả lời ngắn sang dạng Đúng - Sai (Tổng Đúng-Sai 4,0 điểm theo chú thích 3 Phụ lục 1)."}
               - Phần IV. Tự luận: 3,0 điểm (30%).
            2. Tỉ lệ mức độ nhận thức:
               - Biết: 40% (4,0 điểm) | Hiểu: 30% (3,0 điểm) | Vận dụng: 30% (3,0 điểm)

            XUẤT RA ĐẦY ĐỦ 4 PHẦN:
            I. KHUNG MA TRẬN ĐỀ KIỂM TRA (Vẽ bảng Markdown theo Phụ lục 1 CV 7991).
            II. BẢN ĐẶC TẢ ĐỀ KIỂM TRA (Vẽ bảng theo Phụ lục 2 CV 7991).
            III. ĐỀ KIỂM TRA HOÀN CHỈNH.
            IV. ĐÁP ÁN & HƯỚNG DẪN CHẤM (Barem Đúng - Sai chuẩn: đúng 1 ý=0.1đ; 2 ý=0.25đ; 3 ý=0.5đ; 4 ý=1.0đ; barem tự luận chi tiết).
            """
            st.session_state["out_exam"] = call_gemini(prompt_ex, api_key)

    if st.session_state.get("out_exam"):
        st.markdown(st.session_state["out_exam"])
        f_exam_docx = export_docx(f"De_Kiem_Tra_7991_{ten_bai_hoc}", st.session_state["out_exam"], gv_name)
        st.download_button(
            "📥 Tải Bộ hồ sơ đề kiểm tra CV 7991 (.docx)", 
            data=f_exam_docx, 
            file_name=f"De_Kiem_Tra_7991_{ten_bai_hoc}_{gv_name}.docx", 
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

# TAB 3: TRỢ LÝ SÁNG KIẾN KINH NGHIỆM (BÀI TẬP 2)
with tab3:
    st.subheader("💡 Trợ lý Xây dựng Sáng kiến kinh nghiệm (SKKN)")
    ten_skkn_gv = st.text_input(
        "Tên đề tài SKKN:", 
        value="Biện pháp tích hợp Trợ lý Trí tuệ nhân tạo (AI) và Sơ đồ tư duy số hóa nhằm nâng cao năng lực đọc hiểu truyện đồng thoại cho học sinh lớp 6 theo Chương trình GDPT 2018"
    )
    thuc_trang_gv = st.text_area(
        "Thực trạng trước khi áp dụng:", 
        value="Học sinh lớp 6 tiếp cận truyện đồng thoại còn thụ động, chỉ tóm tắt cốt truyện đơn thuần mà chưa phân tích sâu tính cách nhân vật và bài học đường đời; kĩ năng viết kết nối với đọc còn khô cứng."
    )
    if st.button("🚀 Khởi tạo Dự thảo SKKN hoàn chỉnh", type="primary", key="btn_run_skkn"):
        with st.spinner("Đang biên soạn bản dự thảo SKKN theo Thông tư 18/2013/TT-BKHCN..."):
            prompt_sk = f"""
            Chấp bút một bản DỰ THẢO SÁNG KIẾN KINH NGHIỆM hoàn chỉnh theo thể thức chuẩn quy định của ngành Giáo dục:
            - Tên sáng kiến: {ten_skkn_gv}
            - Tác giả: {gv_name} | Đơn vị: Trường THCS Nguyễn Đáng | Môn: {mon_hoc} lớp 6
            - Thực trạng: {thuc_trang_gv}

            CẤU TRÚC ĐẦY ĐỦ:
            1. ĐƠN YÊU CẦU CÔNG NHẬN SÁNG KIẾN (Theo mẫu Phụ lục I Thông tư 18/2013/TT-BKHCN).
            2. PHẦN MỞ ĐẦU (Lý do chọn đề tài, Mục đích, Đối tượng và phạm vi nghiên cứu).
            3. PHẦN NỘI DUNG:
               - Cơ sở lý luận và thực tiễn.
               - Các biện pháp thực hiện (Trình bày chi tiết 3 biện pháp then chốt kèm ví dụ minh họa bài dạy 'Bài học đường đời đầu tiên').
               - Tính mới và tính sáng tạo của sáng kiến.
               - Hiệu quả mang lại (Có bảng số liệu so sánh đối chứng trước và sau khi áp dụng).
            4. KẾT LUẬN VÀ KIẾN NGHỊ (Bài học kinh nghiệm, đề xuất với nhà trường và Phòng/Sở GD&ĐT).
            """
            st.session_state["out_skkn"] = call_gemini(prompt_sk, api_key)

    if st.session_state.get("out_skkn"):
        st.markdown(st.session_state["out_skkn"])
        f_skkn_docx = export_docx(f"SKKN_{ten_skkn_gv[:30]}", st.session_state["out_skkn"], gv_name)
        st.download_button(
            "📥 Tải về Dự thảo SKKN (.docx)", 
            data=f_skkn_docx, 
            file_name=f"Du_Thao_SKKN_{gv_name}.docx", 
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

# TAB 4: CHUYỂN KHBD THÀNH SLIDE POWERPOINT (BÀI TẬP 3 CHUYÊN SÂU - CÓ NẠP FILE)
with tab4:
    st.subheader("🖥️ Trợ lý Nạp Kế hoạch bài dạy & Chuyển đổi thành Bộ Slide (.PPTX)")
    st.caption("Tiếp nhận tệp KHBD (.docx/.txt), phân tích cấu trúc và thiết kế slide phân cấp TO - VỪA - NHỎ kèm Lời dẫn giáo viên")
    
    source_mode = st.radio(
        "Chọn phương thức nạp Kế hoạch bài dạy:", 
        ["📁 Tải lên tệp KHBD (.docx, .txt)", "✏️ Dán nội dung văn bản KHBD", "🔄 Sử dụng KHBD vừa tạo ở Tab 1"], 
        horizontal=True
    )
    
    final_khbd_content = ""
    
    if source_mode == "📁 Tải lên tệp KHBD (.docx, .txt)":
        uploaded_doc = st.file_uploader("Chọn file Kế hoạch bài dạy từ máy tính của cô:", type=["docx", "txt"], key="uploader_khbd_yen")
        if uploaded_doc is not None:
            with st.spinner("Đang trích xuất nội dung văn bản từ tệp..."):
                final_khbd_content = read_uploaded_file(uploaded_doc)
            if final_khbd_content and final_khbd_content.strip():
                st.success(f"✅ Đã nạp thành công tệp: **{uploaded_doc.name}** ({len(final_khbd_content)} ký tự)")
                with st.expander("👁️ Xem trước nội dung đã trích xuất từ tệp"):
                    st.text_area("Nội dung file:", value=final_khbd_content, height=180, disabled=True)
            else:
                st.warning("⚠️ Không thể đọc được nội dung văn bản từ file này. Vui lòng chuyển sang tab '✏️ Dán nội dung văn bản KHBD' để dán trực tiếp!")
                
    elif source_mode == "✏️ Dán nội dung văn bản KHBD":
        final_khbd_content = st.text_area(
            "Dán toàn bộ nội dung giáo án / Kế hoạch bài dạy vào đây:",
            height=220,
            value=f"KẾ HOẠCH BÀI DẠY: {ten_bai_hoc}\nMôn học: {mon_hoc} - {khoi_lop}\nThời lượng: {thoi_luong} tiết\nYêu cầu cần đạt: {yeu_cau_can_dat}"
        )
        
    elif source_mode == "🔄 Sử dụng KHBD vừa tạo ở Tab 1":
        final_khbd_content = st.session_state.get("out_khbd", "")
        if final_khbd_content:
            st.success("✅ Đã lấy thành công Kế hoạch bài dạy từ Tab 1.")
            with st.expander("👁️ Xem lại nội dung KHBD từ Tab 1"):
                st.markdown(final_khbd_content)
        else:
            st.info("💡 Bạn chưa bấm tạo ở Tab 1. Hãy bấm sang mục Dán văn bản hoặc Tải file lên.")

    col_s1, col_s2 = st.columns(2)
    with col_s1:
        so_slide_target = st.slider("Số lượng Slide mục tiêu:", 6, 16, 8)
    with col_s2:
        kieu_thiet_ke = st.selectbox("Phong cách sư phạm:", ["Chuẩn hóa bám sát tiến trình KHBD", "Tương tác phát vấn & Hoạt động nhóm", "Trực quan hóa trọng tâm"])

    if st.button("🚀 Bắt đầu chuyển đổi KHBD thành Bộ Slide PPTX", type="primary", key="btn_run_slide_bt3"):
        if not final_khbd_content or not final_khbd_content.strip():
            st.error("⚠️ Chưa có nội dung Kế hoạch bài dạy! Vui lòng tải file hoặc chọn mục 'Dán nội dung' để tiếp tục.")
        else:
            with st.spinner("AI đang phân tích tiến trình bài dạy và thiết kế kịch bản trình chiếu phân cấp..."):
                prompt_slide = f"""
                ĐÓNG VAI TRÒ: Trợ lý thiết kế nội dung slide dạy học THCS–THPT Việt Nam chuyên nghiệp.
                NHIỆM VỤ: Chuyển toàn bộ Kế hoạch bài dạy dưới đây thành bộ slide trình chiếu PowerPoint gồm {so_slide_target} slide.
                GIÁO VIÊN: {gv_name} - GV
                PHONG CÁCH: {kieu_thiet_ke}

                NỘI DUNG KẾ HOẠCH BÀI DẠY ĐƯỢC CUNG CẤP:
                \"\"\"
                {final_khbd_content}
                \"\"\"

                BẮT BUỘC TUÂN THỦ NGHIÊM NGẶT QUY CÁCH SLIDE:
                1. Mỗi slide bắt đầu bằng: "SLIDE: [Tên tiêu đề slide ngắn gọn]"
                2. Nội dung hiển thị trên mặt slide theo cấu trúc phân cấp:
                   - TO: [Câu hỏi, thông điệp hoặc từ khóa trọng tâm - viết hoa/chữ đậm]
                   - VỪA: [3-4 ý chính, nhiệm vụ hoặc bước thực hiện ngắn gọn]
                   - NHỎ: [Gợi ý hình ảnh minh họa tư liệu, nguồn hoặc hướng dẫn phụ]
                3. Dưới mỗi slide BẮT BUỘC có mục: "LỜI DẪN GV: [2-3 câu ngắn gọn giáo viên nói khi giảng slide này, không hiển thị trên mặt slide]"
                4. Tách biệt hoàn toàn phần chữ cho học sinh và lời dẫn cho giáo viên. Không đưa đáp án bài tập trực tiếp lên slide giao việc.
                """
                st.session_state["out_slide_bt3"] = call_gemini(prompt_slide, api_key)

    if st.session_state.get("out_slide_bt3"):
        st.markdown(st.session_state["out_slide_bt3"])
        f_pptx_bt3 = export_advanced_pptx(ten_bai_hoc, st.session_state["out_slide_bt3"], gv_name)
        st.download_button(
            "📥 Tải về Bộ Slide PowerPoint (.pptx)", 
            data=f_pptx_bt3, 
            file_name=f"Slide_{ten_bai_hoc}_{gv_name}.pptx", 
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )
